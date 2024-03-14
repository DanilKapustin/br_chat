import csv
from datetime import datetime
from io import StringIO
from logging import getLogger
from os import walk, path, environ
from zipfile import ZipFile

import aiofiles
import docx2txt
import pdfplumber
from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from chatbot.db.model.source import Source
from chatbot.dto import KnowledgeDocument
from chatbot.service import knowledge as knowledge_service

FILE_STORAGE_PATH = environ.get("FILE_STORAGE_PATH", "/tmp")
FILE_READ_CHUNK = 1024
ALLOWED_FILE_EXTENSIONS = ["txt", "csv", "md", "pdf", "docx", "xlsx", "zip", "xml"]

logger = getLogger(__name__)


async def save_file(
    db: AsyncSession, source: Source, uploaded_file: UploadFile
) -> Source:
    """Save uploaded file"""
    logger.debug("save_file, source=%s, uploaded_file=%s", source, uploaded_file)
    extension: str = uploaded_file.filename.split(".")[-1]

    async with aiofiles.tempfile.NamedTemporaryFile(
        "w+b", dir=FILE_STORAGE_PATH, delete=False, suffix=f".{extension}"
    ) as out_file:
        while content := await uploaded_file.read(FILE_READ_CHUNK):
            await out_file.write(content)

        source.progress.temporary_file_path = out_file.name
        source.updated_by = "admin"
        source.updated_at = datetime.now()

    await db.commit()
    await db.refresh(source)

    return source


async def _extract_from_txt(
    file_name: str, file_path: str, url: str, extension: str
) -> list[KnowledgeDocument]:
    """Extract text from text file"""
    logger.debug(
        "_extract_from_txt, file_name=%s, file_path=%s, url=%s, extension=%s",
        file_name,
        file_path,
        url,
        extension,
    )

    with open(file_path, "rt", encoding="utf-8", errors="ignore") as f:
        text = f.read()

    if text is None:
        raise RuntimeError("text is empty")

    return [KnowledgeDocument(title=file_name, type=extension, text=text, url=url)]


async def _extract_from_pdf(
    file_name: str, file_path: str, url: str
) -> list[KnowledgeDocument]:
    """Extract text from pdf file"""
    logger.debug(
        "_extract_from_pdf, file_name=%s, file_path=%s, url=%s",
        file_name,
        file_path,
        url,
    )
    file_text: str = ""

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text(x_tolerance=1)

            if page_text is not None:
                file_text += page_text

    return [KnowledgeDocument(title=file_name, type="pdf", text=file_text, url=url)]


async def _extract_from_docx(
    file_name: str, file_path: str, url: str
) -> list[KnowledgeDocument]:
    """Extract text from docx file"""
    logger.debug(
        "_extract_from_docx, file_name=%s, file_path=%s, url=%s",
        file_name,
        file_path,
        url,
    )
    return [
        KnowledgeDocument(
            title=file_name, type="docx", text=docx2txt.process(file_path), url=url
        )
    ]


async def _extract_from_xlsx(
    file_name: str, file_path: str, url: str
) -> list[KnowledgeDocument]:
    """Extract text from xlsx file"""
    logger.debug(
        "_extract_from_xlsx, file_name=%s, file_path=%s, url=%s",
        file_name,
        file_path,
        url,
    )
    workbook: Workbook = load_workbook(file_path)
    result: list[KnowledgeDocument] = []

    for worksheet in workbook:
        csv_output = StringIO()
        writer = csv.writer(csv_output)

        for row in worksheet.values:
            writer.writerow(row)

        result.append(
            KnowledgeDocument(
                title=file_name,
                subtitle=worksheet.title,
                type="csv",
                text=csv_output.getvalue(),
                url=url,
            )
        )

    return result


async def _extract_from_zip(
    file_name: str,
    file_path: str,
    url: str,
) -> list[KnowledgeDocument]:
    """Extract text from zip file"""
    logger.debug(
        "_extract_from_zip, file_name=%s, file_path=%s, url=%s",
        file_name,
        file_path,
        url,
    )
    result: list[KnowledgeDocument] = []

    async with aiofiles.tempfile.TemporaryDirectory(dir=FILE_STORAGE_PATH) as extracted:
        with ZipFile(file_path) as zip_file:
            logger.debug("_extract_from_zip, extracting to=%s", extracted)
            zip_file.extractall(extracted)
            ctr: int = 0

            for root, _, file_names in walk(extracted):
                ctr += 1
                logger.debug(
                    "_extract_from_zip, root=%s, file_names=%s", root, file_names
                )

                file_names = filter(
                    lambda x: not x.startswith(".")
                    and x.split(".")[-1] in ALLOWED_FILE_EXTENSIONS,
                    file_names,
                )

                for inner_file_name in file_names:
                    logger.debug(
                        "_extract_from_zip, inner_file_name=%s", inner_file_name
                    )

                    try:
                        result += await _extract_text(
                            inner_file_name,
                            path.join(root, inner_file_name),
                            path.join(url, root.replace(extracted, "")),
                        )

                    except NotImplementedError:
                        logger.warning(
                            "_extract_from_zip, file=%s cannot be processed",
                            inner_file_name,
                        )

            if len(result) == 0:
                raise RuntimeError("Archive is empty")

    return result


async def _extract_text(
    file_name: str, file_path: str, url: str
) -> list[KnowledgeDocument]:
    """Extract text from file(s)"""
    # depending on file extension, use different function
    logger.debug(
        "_extract_text, file_name=%s, file_path=%s, url=%s", file_name, file_path, url
    )
    extension: str = file_path.split(".")[-1]
    url = path.join(url, file_name)

    match extension:
        case "txt" | "csv" | "md" | "xml":
            return await _extract_from_txt(file_name, file_path, url, extension)
        case "pdf":
            return await _extract_from_pdf(file_name, file_path, url)
        case "docx":
            return await _extract_from_docx(file_name, file_path, url)
        case "xlsx":
            return await _extract_from_xlsx(file_name, file_path, url)
        case "zip":
            return await _extract_from_zip(file_name, file_path, url)

    raise NotImplementedError()


async def index(source: Source) -> int:
    """Index source"""
    logger.info("index, source=%s", source)
    documents: list[KnowledgeDocument] = await _extract_text(
        source.title, source.progress.temporary_file_path, ""
    )

    for document in documents:
        logger.debug("index, document=%s", document)
        await knowledge_service.create(str(source.id), source.title, document)

    return len(documents)
